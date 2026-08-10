"""Serialize integrated GT-CCS results without coupling to Streamlit.

The export layer accepts the mapping returned by ``ScenarioRunner.run`` and
produces either a strict JSON document or an in-memory Excel workbook. It does
not recalculate, round, or otherwise modify engineering and economic results.
"""

from __future__ import annotations

from collections.abc import Mapping, Sequence
from dataclasses import asdict, is_dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
import json
import math
from pathlib import Path
import re
from typing import Any

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill


EXPORT_SCHEMA_VERSION = "1.0"

RESULT_TABLE_SHEETS = {
    "annual_technical_ledger": "annual_technical",
    "annual_cash_flow_df": "annual_cash_flow",
    "mbal_df": "storage",
    "vfp_co2_df": "co2_well",
    "vfp_gt_df": "geothermal",
    "well_pressure_df": "well_pressures",
    "doublet_df": "thermal_front",
    "relative_permeability_df": "relative_permeability",
    "annual_relative_permeability_df": "annual_rel_perm",
    "annual_injection_diagnostics_df": "annual_injection_debug",
    "co2_price_df": "co2_price",
}

SUMMARY_RESULT_KEYS = (
    "operational_kpis",
    "economic_kpis",
    "gt_info",
    "scenario_notes",
    "transport_configuration",
)

_INVALID_SHEET_CHARACTERS = re.compile(r"[\\/*?:\[\]]")
_HEADER_FILL = "D9EAF7"
_HEADER_FONT = "17365D"


def _json_safe(value: Any) -> Any:
    """Return a value accepted by ``json.dumps(..., allow_nan=False)``."""

    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if is_dataclass(value) and not isinstance(value, type):
        return _json_safe(asdict(value))
    if isinstance(value, pd.DataFrame):
        records = []
        for record in value.to_dict(orient="records"):
            records.append(
                {str(key): _json_safe(item) for key, item in record.items()}
            )
        return {
            "columns": [str(column) for column in value.columns],
            "records": records,
            "attrs": _json_safe(dict(value.attrs)),
        }
    if isinstance(value, pd.Series):
        return [_json_safe(item) for item in value.tolist()]
    if isinstance(value, Mapping):
        return {str(key): _json_safe(item) for key, item in value.items()}
    if isinstance(value, np.ndarray):
        return [_json_safe(item) for item in value.tolist()]
    if isinstance(value, np.generic):
        return _json_safe(value.item())
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, Path):
        return value.as_posix()
    if isinstance(value, Decimal):
        value = float(value)
    if isinstance(value, float):
        return value if math.isfinite(value) else None
    if isinstance(value, (str, int, bool)):
        return value
    if isinstance(value, Sequence) and not isinstance(
        value, (str, bytes, bytearray)
    ):
        return [_json_safe(item) for item in value]
    raise TypeError(
        "Rezultat sadrži vrijednost koja se ne može izvesti u JSON: "
        f"{type(value).__name__}."
    )


def build_results_json(results: Mapping[str, Any], active_inputs=None, *, indent=2):
    """Build a strict UTF-8 JSON representation of inputs and results."""

    if not isinstance(results, Mapping):
        raise TypeError("results mora biti mapiranje rezultata ScenarioRunnera.")
    payload = {
        "schema_version": EXPORT_SCHEMA_VERSION,
        "active_inputs": _json_safe(active_inputs or {}),
        "results": _json_safe(results),
    }
    return json.dumps(
        payload,
        ensure_ascii=False,
        indent=indent,
        allow_nan=False,
    )


def _excel_scalar(value: Any) -> Any:
    """Normalize one scalar while preventing accidental Excel formulas."""

    normalized = _json_safe(value)
    if isinstance(normalized, (dict, list)):
        normalized = json.dumps(normalized, ensure_ascii=False, allow_nan=False)
    if isinstance(normalized, str) and normalized.startswith(("=", "+", "-", "@")):
        return "'" + normalized
    return normalized


def _excel_frame(frame: pd.DataFrame) -> pd.DataFrame:
    safe_frame = frame.copy()
    safe_frame.columns = [str(column) for column in safe_frame.columns]
    for column in safe_frame.columns:
        safe_frame[column] = safe_frame[column].map(_excel_scalar)
    return safe_frame


def _input_frame(active_inputs) -> pd.DataFrame:
    rows = []
    for parameter, entry in (active_inputs or {}).items():
        if isinstance(entry, Sequence) and not isinstance(
            entry, (str, bytes, bytearray)
        ):
            value = entry[0] if len(entry) > 0 else None
            unit = entry[1] if len(entry) > 1 else ""
            description = entry[2] if len(entry) > 2 else ""
        else:
            value = entry
            unit = ""
            description = ""
        rows.append(
            {
                "parameter": str(parameter),
                "value": _excel_scalar(value),
                "unit": _excel_scalar(unit),
                "description": _excel_scalar(description),
            }
        )
    return pd.DataFrame(
        rows,
        columns=["parameter", "value", "unit", "description"],
    )


def _flatten_summary(section: str, value: Any, prefix=""):
    if isinstance(value, Mapping):
        rows = []
        for key, item in value.items():
            path = f"{prefix}.{key}" if prefix else str(key)
            rows.extend(_flatten_summary(section, item, path))
        return rows
    return [
        {
            "section": section,
            "parameter": prefix,
            "value": _excel_scalar(value),
        }
    ]


def _summary_frame(results: Mapping[str, Any]) -> pd.DataFrame:
    rows = []
    for key in SUMMARY_RESULT_KEYS:
        if key in results:
            rows.extend(_flatten_summary(key, results[key]))
    return pd.DataFrame(rows, columns=["section", "parameter", "value"])


def _unique_sheet_name(base_name: str, used_names: set[str]) -> str:
    cleaned = _INVALID_SHEET_CHARACTERS.sub("_", base_name).strip() or "results"
    cleaned = cleaned[:31]
    candidate = cleaned
    suffix_number = 2
    while candidate.casefold() in used_names:
        suffix = f"_{suffix_number}"
        candidate = cleaned[: 31 - len(suffix)] + suffix
        suffix_number += 1
    used_names.add(candidate.casefold())
    return candidate


def _format_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book
    workbook.properties.title = "GT-CCS integrirani rezultati"
    workbook.properties.subject = "Tehnički i ekonomski rezultati scenarija"
    workbook.properties.creator = "CCUS_cluster"

    header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
    header_font = Font(bold=True, color=_HEADER_FONT)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        if worksheet.max_row > 1 and worksheet.max_column > 0:
            worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column_cells in worksheet.columns:
            values = [
                "" if cell.value is None else str(cell.value)
                for cell in column_cells[: min(worksheet.max_row, 200)]
            ]
            width = min(max((len(value) for value in values), default=8) + 2, 60)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)


def build_results_excel(results: Mapping[str, Any], active_inputs=None) -> bytes:
    """Build an XLSX workbook containing inputs, summaries and all result tables."""

    if not isinstance(results, Mapping):
        raise TypeError("results mora biti mapiranje rezultata ScenarioRunnera.")

    buffer = BytesIO()
    used_names: set[str] = set()
    with pd.ExcelWriter(buffer, engine="openpyxl") as writer:
        _input_frame(active_inputs).to_excel(
            writer,
            sheet_name=_unique_sheet_name("inputs", used_names),
            index=False,
        )
        _summary_frame(results).to_excel(
            writer,
            sheet_name=_unique_sheet_name("summary", used_names),
            index=False,
        )
        for key, value in results.items():
            if not isinstance(value, pd.DataFrame):
                continue
            preferred_name = RESULT_TABLE_SHEETS.get(key, str(key))
            _excel_frame(value).to_excel(
                writer,
                sheet_name=_unique_sheet_name(preferred_name, used_names),
                index=False,
            )
        _format_workbook(writer)
    return buffer.getvalue()


__all__ = [
    "EXPORT_SCHEMA_VERSION",
    "RESULT_TABLE_SHEETS",
    "build_results_excel",
    "build_results_json",
]
